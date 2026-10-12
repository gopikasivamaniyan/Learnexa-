import io
import os
import datetime as dt
from functools import wraps

import jwt
from bson import ObjectId
from dotenv import load_dotenv
from flask import Flask, jsonify, request, g
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash

from llm import ask_ai, ask_ai_json, LLMError

load_dotenv()
app = Flask(__name__)
CORS(app)  # frontend (localhost:8080 / 5173) -> backend (localhost:5000)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB upload limit (matches the UI)

if os.getenv("MONGO_URI") == "mock":
    import mongomock
    db = mongomock.MongoClient()[os.getenv("MONGO_DB", "learnexa")]
else:
    from pymongo import MongoClient
    db = MongoClient(os.getenv("MONGO_URI", "mongodb://localhost:27017"))[os.getenv("MONGO_DB", "learnexa")]

SECRET = os.getenv("JWT_SECRET", "dev-secret-change-me")


# ============================================================ helpers ====

def public_user(u):
    return {"id": str(u["_id"]), "name": u["name"], "email": u["email"], "role": u["role"]}


def make_token(u):
    payload = {"sub": str(u["_id"]), "role": u["role"],
               "exp": dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=7)}
    return jwt.encode(payload, SECRET, algorithm="HS256")


def token_required(fn):
    @wraps(fn)
    def wrapper(*a, **kw):
        h = request.headers.get("Authorization", "")
        if not h.startswith("Bearer "):
            return jsonify(message="Please log in to continue."), 401
        try:
            data = jwt.decode(h[7:], SECRET, algorithms=["HS256"])
            g.user = db.users.find_one({"_id": ObjectId(data["sub"])})
        except Exception:
            return jsonify(message="Session expired. Please log in again."), 401
        if not g.user:
            return jsonify(message="Account not found."), 401
        return fn(*a, **kw)
    return wrapper


def institution_required(fn):
    @wraps(fn)
    @token_required
    def wrapper(*a, **kw):
        if g.user["role"] != "institution":
            return jsonify(message="This action is only available to institution accounts."), 403
        return fn(*a, **kw)
    return wrapper


def oid_or_none(id_str):
    try:
        return ObjectId(id_str)
    except Exception:
        return None


def aware(d):
    """Mongo may return naive datetimes; always treat them as UTC."""
    if isinstance(d, dt.datetime) and d.tzinfo is None:
        return d.replace(tzinfo=dt.timezone.utc)
    return d


def doc_id(d, extra=None):
    out = {"id": str(d["_id"])}
    for k, v in d.items():
        if k == "_id":
            continue
        if isinstance(v, ObjectId):
            v = str(v)
        if isinstance(v, dt.datetime):
            v = v.isoformat()
        out[k] = v
    if extra:
        out.update(extra)
    return out


def owner_filter():
    return {"user_id": g.user["_id"]}


_ocr_engine = None


def ocr_image_bytes(raw):
    """Read text from an image. Uses RapidOCR (pip only); falls back to pytesseract if installed."""
    global _ocr_engine
    try:
        import numpy as np
        from PIL import Image
        img = Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception:
        return ""
    try:
        from rapidocr_onnxruntime import RapidOCR
        if _ocr_engine is None:
            _ocr_engine = RapidOCR()
        result, _ = _ocr_engine(np.array(img))
        if result:
            return "\n".join(line[1] for line in result).strip()
        return ""
    except ImportError:
        pass
    except Exception:
        return ""
    try:
        import pytesseract
        return pytesseract.image_to_string(img).strip()
    except Exception:
        return ""


def ocr_scanned_pdf(raw, max_pages=10):
    """OCR a PDF that has no text layer (needs: pip install pymupdf)."""
    try:
        import fitz  # pymupdf
    except ImportError:
        return ""
    out = []
    try:
        pdf = fitz.open(stream=raw, filetype="pdf")
        for i, page in enumerate(pdf):
            if i >= max_pages:
                break
            png = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
            out.append(ocr_image_bytes(png))
    except Exception:
        pass
    return "\n".join(t for t in out if t)


def safe_json(prompt):
    """ask_ai_json with one automatic retry using a stricter instruction."""
    try:
        return ask_ai_json(prompt)
    except LLMError as first:
        strict = (prompt + "\n\nIMPORTANT: Reply with ONLY valid JSON. Use double quotes for all "
                  "keys and strings. No markdown, no code fences, no explanation, no trailing commas.")
        try:
            return ask_ai_json(strict)
        except LLMError:
            raise first


def as_text(v):
    """AI models sometimes return lists/objects where text is expected. Always give back a string."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, (list, tuple)):
        return "\n".join(("- " + as_text(x)) if not isinstance(x, (list, dict)) else as_text(x) for x in v)
    if isinstance(v, dict):
        return "\n".join(f"{k}: {as_text(x)}" for k, x in v.items())
    return str(v)


def _priority(v):
    p = as_text(v).strip().capitalize()
    return p if p in ("High", "Medium", "Low") else "Medium"


def _answer_index(v, n_options=4):
    """Model may return 2, "2", "C" ... always give a valid 0-based index."""
    try:
        i = int(v)
    except (TypeError, ValueError):
        letter = as_text(v).strip().upper()[:1]
        i = "ABCD".find(letter) if letter else 0
    return max(0, min(i, max(n_options - 1, 0)))


def clean_fields(row, *keys):
    for k in keys:
        row[k] = as_text(row.get(k))
    return row


def extract_text(filename, raw):
    """Pull readable text out of an uploaded file (PDF, DOCX, plain text)."""
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    text = ""
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
            if len(text.strip()) < 20:  # scanned PDF -> OCR
                text = ocr_scanned_pdf(raw) or text
        elif ext == "docx":
            try:
                import docx  # pip install python-docx (optional)
                d = docx.Document(io.BytesIO(raw))
                parts = [p.text for p in d.paragraphs]
                for table in d.tables:
                    for row in table.rows:
                        parts.append(" | ".join(c.text.strip() for c in row.cells))
                text = "\n".join(x for x in parts if x.strip())
            except ImportError:
                text = ""
        elif ext in ("png", "jpg", "jpeg", "gif", "webp", "bmp"):
            text = ocr_image_bytes(raw)
        else:
            text = raw.decode("utf-8", errors="ignore")
    except Exception:
        text = ""
    return text[:20000]


# ================================================================ auth ====

@app.post("/api/auth/register/learner")
def register_learner():
    d = request.get_json(silent=True) or {}
    name, email, pw = (d.get("name") or "").strip(), (d.get("email") or "").strip().lower(), d.get("password") or ""
    if not name or not email or len(pw) < 6:
        return jsonify(message="Name, email and a password of 6+ characters are required."), 400
    if db.users.find_one({"email": email}):
        return jsonify(message="An account with this email already exists."), 409
    u = {"name": name, "email": email, "role": "learner", "level": d.get("level", ""),
         "goal": d.get("goal", ""), "password": generate_password_hash(pw),
         "streak": 0, "last_active": None,
         "created_at": dt.datetime.now(dt.timezone.utc)}
    u["_id"] = db.users.insert_one(u).inserted_id
    return jsonify(token=make_token(u), user=public_user(u)), 201


@app.post("/api/auth/register/institution")
def register_institution():
    d = request.get_json(silent=True) or {}
    name = (d.get("institutionName") or d.get("name") or "").strip()
    email, pw = (d.get("email") or "").strip().lower(), d.get("password") or ""
    if not name or not email or len(pw) < 6:
        return jsonify(message="Institution name, email and a password of 6+ characters are required."), 400
    if db.users.find_one({"email": email}):
        return jsonify(message="An account with this email already exists."), 409
    u = {"name": name, "email": email, "role": "institution",
         "contactPerson": d.get("contactPerson", ""), "phone": d.get("phone", ""),
         "password": generate_password_hash(pw),
         "created_at": dt.datetime.now(dt.timezone.utc)}
    u["_id"] = db.users.insert_one(u).inserted_id
    return jsonify(token=make_token(u), user=public_user(u)), 201


@app.post("/api/auth/login")
def login():
    d = request.get_json(silent=True) or {}
    u = db.users.find_one({"email": (d.get("email") or "").strip().lower()})
    if not u or not check_password_hash(u["password"], d.get("password") or ""):
        return jsonify(message="Incorrect email or password."), 401
    if d.get("role") and d["role"] != u["role"]:
        return jsonify(message=f"This is a {u['role']} account. Switch the tab above and try again."), 403
    return jsonify(token=make_token(u), user=public_user(u))


@app.get("/api/auth/me")
@token_required
def me():
    return jsonify(user=public_user(g.user))


# ============================================================== streak ====

def touch_streak():
    today = dt.datetime.now(dt.timezone.utc).date()
    last = g.user.get("last_active")
    streak = g.user.get("streak", 0)
    if last:
        last = last.date() if isinstance(last, dt.datetime) else last
        gap = (today - last).days
        if gap == 1:
            streak += 1
        elif gap > 1:
            streak = 1
        # gap == 0 -> already counted today
    else:
        streak = 1
    db.users.update_one({"_id": g.user["_id"]},
                         {"$set": {"streak": streak, "last_active": dt.datetime.now(dt.timezone.utc)}})
    g.user["streak"] = streak


# =========================================================== AI Tutor ====

def _save_chat_pair(user_text, reply, lang):
    now = dt.datetime.now(dt.timezone.utc)
    db.chats.insert_many([
        {"user_id": g.user["_id"], "role": "user", "text": user_text, "language": lang, "created_at": now},
        {"user_id": g.user["_id"], "role": "assistant", "text": reply, "language": lang,
         "created_at": now + dt.timedelta(milliseconds=1)},
    ])


@app.post("/api/tutor/chat")
@token_required
def tutor_chat():
    d = request.get_json(silent=True) or {}
    message = (d.get("message") or "").strip()
    if not message:
        return jsonify(message="Message is empty."), 400
    lang = d.get("language", "en")
    past = list(db.chats.find(owner_filter()).sort("created_at", -1).limit(8))
    history = [{"role": m["role"], "content": m["text"]} for m in reversed(past)]
    try:
        reply = ask_ai(message, history, lang)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    _save_chat_pair(message, reply, lang)
    touch_streak()
    return jsonify(reply=reply)


@app.get("/api/tutor/history")
@token_required
def tutor_history():
    rows = list(db.chats.find(owner_filter()).sort("created_at", 1))
    # Group into "sessions": a new session starts after a >30min gap.
    sessions = []
    cur = None
    for m in rows:
        if m["role"] != "user":
            continue
        if cur and (aware(m["created_at"]) - aware(cur["last_at"])).total_seconds() < 1800:
            cur["count"] += 1
            cur["last_at"] = m["created_at"]
        else:
            if cur:
                sessions.append(cur)
            cur = {"id": str(m["_id"]), "title": m["text"][:60], "preview": m["text"][:80],
                   "date": m["created_at"].date().isoformat(), "count": 1, "last_at": m["created_at"]}
    if cur:
        sessions.append(cur)
    sessions.reverse()
    out = [{"id": s["id"], "title": s["title"], "preview": s["preview"],
            "date": s["date"], "messages": s["count"]} for s in sessions]
    return jsonify(sessions=out)


@app.delete("/api/tutor/history")
@token_required
def tutor_clear():
    db.chats.delete_many(owner_filter())
    return "", 204


@app.delete("/api/tutor/history/<session_id>")
@token_required
def tutor_delete_session(session_id):
    anchor = db.chats.find_one({"_id": oid_or_none(session_id), **owner_filter()})
    if not anchor:
        return jsonify(message="Conversation not found."), 404
    # Delete every message in the same 30-minute session window as the anchor message.
    rows = list(db.chats.find(owner_filter()).sort("created_at", 1))
    window_ids = []
    in_session, last_t = False, None
    for m in rows:
        t = aware(m["created_at"])
        gap_new = last_t is None or (t - last_t).total_seconds() >= 1800
        if gap_new:
            in_session = False
        if m["_id"] == anchor["_id"]:
            in_session = True
        if in_session:
            window_ids.append(m["_id"])
        last_t = t
    db.chats.delete_many({"_id": {"$in": window_ids}, **owner_filter()})
    return "", 204


@app.post("/api/tutor/image")
@token_required
def tutor_image():
    question = (request.form.get("question") or "").strip()
    lang = request.form.get("language", "en")
    image = request.files.get("image")
    if not question:
        return jsonify(message="Please add a question about the image."), 400
    prompt = question
    if image:
        ocr_text = ocr_image_bytes(image.read())
        if ocr_text:
            prompt = (f"The learner attached an image. Text extracted from the image by OCR:\n"
                       f"\"\"\"\n{ocr_text[:4000]}\n\"\"\"\n\n"
                       f"Using this text, answer the learner's question: {question}")
        else:
            prompt = (f"The learner attached an image named '{image.filename}', but no readable text "
                       "could be extracted from it (it may be a photo or diagram). Answer using the "
                       f"learner's question and general subject knowledge. Question: {question}")
    try:
        reply = ask_ai(prompt, None, lang)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    _save_chat_pair(question, reply, lang)
    touch_streak()
    return jsonify(reply=reply)


# ============================================================ Documents ====

@app.get("/api/documents")
@token_required
def list_documents():
    rows = db.documents.find(owner_filter()).sort("uploaded_at", -1)
    return jsonify(documents=[doc_id(r, {"uploadedAt": r["uploaded_at"].date().isoformat()}) for r in rows])


@app.post("/api/documents")
@token_required
def upload_document():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(message="No file received."), 400
    ext = (f.filename.rsplit(".", 1)[-1] if "." in f.filename else "").lower()
    ftype = ("pdf" if ext == "pdf" else "docx" if ext in ("docx", "doc")
             else "image" if ext in ("png", "jpg", "jpeg", "gif", "webp", "bmp") else "pdf")
    raw = f.read()
    size_kb = len(raw) / 1024
    size = f"{size_kb/1024:.1f} MB" if size_kb > 1024 else f"{size_kb:.0f} KB"
    text = extract_text(f.filename, raw)
    d = {"user_id": g.user["_id"], "name": f.filename, "type": ftype, "size": size,
         "pages": max(1, len(text) // 2000) if text else 1,
         "subject": request.form.get("subject", "General"), "status": "ready",
         "text": text, "uploaded_at": dt.datetime.now(dt.timezone.utc)}
    d["_id"] = db.documents.insert_one(d).inserted_id
    return jsonify(document=doc_id(d, {"uploadedAt": d["uploaded_at"].date().isoformat()})), 201


@app.delete("/api/documents/<id>")
@token_required
def delete_document(id):
    db.documents.delete_one({"_id": oid_or_none(id), **owner_filter()})
    return "", 204


def _get_owned_doc(id):
    doc = db.documents.find_one({"_id": oid_or_none(id), **owner_filter()})
    return doc


DOC_CTX_CHARS = int(os.getenv("DOC_CTX_CHARS", "6000"))


def _doc_context(doc):
    if doc.get("text"):
        return f"Document '{doc['name']}' (subject: {doc['subject']}) content:\n{doc['text'][:DOC_CTX_CHARS]}"
    return f"Document '{doc['name']}' (subject: {doc['subject']}). No readable text could be extracted from this file."


@app.post("/api/documents/<id>/chat")
@token_required
def document_chat(id):
    doc = _get_owned_doc(id)
    if not doc:
        return jsonify(message="Document not found."), 404
    d = request.get_json(silent=True) or {}
    message = (d.get("message") or "").strip()
    if not message:
        return jsonify(message="Message is empty."), 400
    lang = d.get("language", "en")
    prompt = f"{_doc_context(doc)}\n\nLearner's question about this document: {message}"
    try:
        reply = ask_ai(prompt, None, lang)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    _save_chat_pair(message, reply, lang)
    touch_streak()
    return jsonify(reply=reply)


@app.post("/api/documents/<id>/summary")
@token_required
def document_summary(id):
    doc = _get_owned_doc(id)
    if not doc:
        return jsonify(message="Document not found."), 404
    prompt = f"{_doc_context(doc)}\n\nWrite a clear, well-structured summary of this document for a student, using headings and bullet points."
    try:
        reply = ask_ai(prompt, None, "en")
    except LLMError as e:
        return jsonify(message=str(e)), 502
    return jsonify(summary=reply)


@app.post("/api/documents/<id>/notes")
@token_required
def document_notes(id):
    doc = _get_owned_doc(id)
    if not doc:
        return jsonify(message="Document not found."), 404
    prompt = (f"{_doc_context(doc)}\n\nWrite study notes for this document as a JSON object: "
              '{"title": str, "content": str}. content should use \\n for line breaks and bullet points.')
    try:
        data = safe_json(prompt)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    if isinstance(data, list):
        data = data[0] if data and isinstance(data[0], dict) else {}
    if not isinstance(data, dict):
        data = {}
    note = {"user_id": g.user["_id"], "title": as_text(data.get("title")) or doc["name"],
            "subject": doc["subject"], "source": "ai", "content": as_text(data.get("content")),
            "updated_at": dt.datetime.now(dt.timezone.utc)}
    note["_id"] = db.notes.insert_one(note).inserted_id
    return jsonify(note=doc_id(note, {"updatedAt": note["updated_at"].date().isoformat()})), 201


@app.post("/api/documents/<id>/flashcards")
@token_required
def document_flashcards(id):
    doc = _get_owned_doc(id)
    if not doc:
        return jsonify(message="Document not found."), 404
    prompt = (f"{_doc_context(doc)}\n\nCreate 6 flashcards from this document as a JSON array of "
              '{"front": str, "back": str}.')
    try:
        data = safe_json(prompt)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    cards = []
    for c in (x for x in (data if isinstance(data, list) else []) if isinstance(x, dict)):
        card = {"user_id": g.user["_id"], "front": as_text(c.get("front")), "back": as_text(c.get("back")),
                "subject": doc["subject"], "status": "new", "created_at": dt.datetime.now(dt.timezone.utc)}
        card["_id"] = db.flashcards.insert_one(card).inserted_id
        cards.append(doc_id(card))
    return jsonify(flashcards=cards), 201


@app.post("/api/documents/<id>/quiz")
@token_required
def document_quiz(id):
    doc = _get_owned_doc(id)
    if not doc:
        return jsonify(message="Document not found."), 404
    return _generate_quiz(doc["subject"], _doc_context(doc))


# ================================================================ Notes ====

@app.get("/api/notes")
@token_required
def list_notes():
    rows = [clean_fields(r, "title", "subject", "content")
            for r in db.notes.find(owner_filter()).sort("updated_at", -1)]
    return jsonify(notes=[doc_id(r, {"updatedAt": r["updated_at"].date().isoformat()}) for r in rows])


@app.post("/api/notes")
@token_required
def create_note():
    d = request.get_json(silent=True) or {}
    note = {"user_id": g.user["_id"], "title": (d.get("title") or "Untitled note").strip(),
            "subject": d.get("subject", "General"), "source": "manual",
            "content": d.get("content", ""), "updated_at": dt.datetime.now(dt.timezone.utc)}
    note["_id"] = db.notes.insert_one(note).inserted_id
    return jsonify(note=doc_id(note, {"updatedAt": note["updated_at"].date().isoformat()})), 201


@app.put("/api/notes/<id>")
@token_required
def update_note(id):
    d = request.get_json(silent=True) or {}
    fields = {k: d[k] for k in ("title", "subject", "content") if k in d}
    fields["updated_at"] = dt.datetime.now(dt.timezone.utc)
    db.notes.update_one({"_id": oid_or_none(id), **owner_filter()}, {"$set": fields})
    row = db.notes.find_one({"_id": oid_or_none(id), **owner_filter()})
    if not row:
        return jsonify(message="Note not found."), 404
    return jsonify(note=doc_id(row, {"updatedAt": row["updated_at"].date().isoformat()}))


@app.delete("/api/notes/<id>")
@token_required
def delete_note(id):
    db.notes.delete_one({"_id": oid_or_none(id), **owner_filter()})
    return "", 204


# =========================================================== Flashcards ====

@app.get("/api/flashcards")
@token_required
def list_flashcards():
    rows = [clean_fields(r, "front", "back", "subject") for r in db.flashcards.find(owner_filter())]
    return jsonify(flashcards=[doc_id(r) for r in rows])


@app.post("/api/flashcards")
@token_required
def create_flashcard():
    d = request.get_json(silent=True) or {}
    if d.get("generate"):
        subject = d.get("subject", "General")
        prompt = f"Create 6 flashcards about {subject} for a student, as a JSON array of {{\"front\": str, \"back\": str}}."
        try:
            data = safe_json(prompt)
        except LLMError as e:
            return jsonify(message=str(e)), 502
        cards = []
        for c in (x for x in (data if isinstance(data, list) else []) if isinstance(x, dict)):
            card = {"user_id": g.user["_id"], "front": as_text(c.get("front")), "back": as_text(c.get("back")),
                     "subject": subject, "status": "new", "created_at": dt.datetime.now(dt.timezone.utc)}
            card["_id"] = db.flashcards.insert_one(card).inserted_id
            cards.append(doc_id(card))
        return jsonify(flashcards=cards), 201
    card = {"user_id": g.user["_id"], "front": d.get("front", ""), "back": d.get("back", ""),
            "subject": d.get("subject", "General"), "status": "new",
            "created_at": dt.datetime.now(dt.timezone.utc)}
    card["_id"] = db.flashcards.insert_one(card).inserted_id
    return jsonify(flashcard=doc_id(card)), 201


@app.patch("/api/flashcards/<id>")
@token_required
def update_flashcard(id):
    d = request.get_json(silent=True) or {}
    fields = {k: d[k] for k in ("status", "front", "back") if k in d}
    db.flashcards.update_one({"_id": oid_or_none(id), **owner_filter()}, {"$set": fields})
    row = db.flashcards.find_one({"_id": oid_or_none(id), **owner_filter()})
    if not row:
        return jsonify(message="Flashcard not found."), 404
    return jsonify(flashcard=doc_id(row))


@app.delete("/api/flashcards/<id>")
@token_required
def delete_flashcard(id):
    db.flashcards.delete_one({"_id": oid_or_none(id), **owner_filter()})
    return "", 204


# ================================================================= Quiz ====

def _generate_quiz(subject, context=None):
    prompt = (context or f"Topic: {subject}.") + (
        '\n\nCreate 5 multiple-choice quiz questions as a JSON array of objects: '
        '{"question": str, "options": [4 strings], "answerIndex": int (0-3), '
        '"topic": str, "explanation": str}.')
    try:
        data = safe_json(prompt)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    questions = []
    for i, q in enumerate(x for x in (data if isinstance(data, list) else []) if isinstance(x, dict)):
        questions.append({
            "id": f"gq{i}", "question": as_text(q.get("question")),
            "options": [as_text(o) for o in (q.get("options") or [])][:4],
            "answerIndex": _answer_index(q.get("answerIndex")), "topic": as_text(q.get("topic")) or subject,
            "explanation": as_text(q.get("explanation")),
        })
    questions = [q for q in questions if q["question"] and len(q["options"]) >= 2]
    for q in questions:
        q["answerIndex"] = min(q["answerIndex"], len(q["options"]) - 1)
    if not questions:
        return jsonify(message="The AI could not make quiz questions. Please try again."), 502
    return jsonify(questions=questions)


@app.get("/api/quiz")
@token_required
def get_quiz():
    subject = request.args.get("subject", "General knowledge")
    return _generate_quiz(subject)


@app.post("/api/quiz/submit")
@token_required
def submit_quiz():
    d = request.get_json(silent=True) or {}
    questions = d.get("questions", [])
    answers = d.get("answers", {})  # {questionId: selectedIndex}
    correct, topics = 0, {}
    for q in questions:
        qid = q.get("id")
        topic = q.get("topic", "General")
        is_correct = answers.get(qid) == q.get("answerIndex")
        topics.setdefault(topic, {"correct": 0, "total": 0})
        topics[topic]["total"] += 1
        if is_correct:
            correct += 1
            topics[topic]["correct"] += 1
    total = len(questions) or 1
    percentage = round(correct / total * 100)
    attempt = {"user_id": g.user["_id"], "total": total, "correct": correct, "percentage": percentage,
               "topics": topics, "created_at": dt.datetime.now(dt.timezone.utc)}
    db.quiz_attempts.insert_one(attempt)
    now = dt.datetime.now(dt.timezone.utc)
    for topic, s in topics.items():
        strength = round(s["correct"] / s["total"] * 100)
        if strength < 70:
            db.revision_items.update_one(
                {"user_id": g.user["_id"], "topic": topic},
                {"$set": {"strength": strength, "due_at": now, "subject": topic}},
                upsert=True)
        else:
            db.revision_items.delete_one({"user_id": g.user["_id"], "topic": topic})
    touch_streak()
    strong = [t for t, s in topics.items() if s["correct"] / s["total"] >= 0.7]
    weak = [t for t, s in topics.items() if s["correct"] / s["total"] < 0.7]
    return jsonify(score=correct, total=total, percentage=percentage,
                   strongTopics=strong, weakTopics=weak)


# ============================================================= Planner ====

@app.get("/api/study-plan")
@token_required
def get_study_plan():
    rows = db.study_plan.find(owner_filter()).sort("created_at", -1)
    return jsonify(tasks=[doc_id(r) for r in rows])


@app.post("/api/study-plan")
@token_required
def create_study_plan():
    d = request.get_json(silent=True) or {}
    prompt = (
        f"Subject: {d.get('subject','General')}. Topics: {d.get('topics','')}. "
        f"Exam date: {d.get('examDate','not specified')}. Available study hours/day: {d.get('hours','1')}. "
        f"Difficulty: {d.get('difficulty','Medium')}. Preferred study time: {d.get('preferredTime','Any')}.\n"
        'Create a daily study plan as a JSON array of {"topic": str, "duration": str, "priority": '
        '"High"|"Medium"|"Low"}. Keep it to 4-8 tasks.'
    )
    try:
        data = safe_json(prompt)
    except LLMError as e:
        return jsonify(message=str(e)), 502
    db.study_plan.delete_many(owner_filter())
    tasks = []
    for t in (x for x in (data if isinstance(data, list) else []) if isinstance(x, dict)):
        task = {"user_id": g.user["_id"], "topic": as_text(t.get("topic")), "duration": as_text(t.get("duration")) or "30 min",
                 "priority": _priority(t.get("priority")), "done": False,
                 "created_at": dt.datetime.now(dt.timezone.utc)}
        task["_id"] = db.study_plan.insert_one(task).inserted_id
        tasks.append(doc_id(task))
    return jsonify(tasks=tasks), 201


@app.patch("/api/study-plan/<id>")
@token_required
def toggle_study_task(id):
    d = request.get_json(silent=True) or {}
    db.study_plan.update_one({"_id": oid_or_none(id), **owner_filter()}, {"$set": {"done": bool(d.get("done"))}})
    row = db.study_plan.find_one({"_id": oid_or_none(id), **owner_filter()})
    if not row:
        return jsonify(message="Task not found."), 404
    return jsonify(task=doc_id(row))


# ============================================================ Revision ====

@app.get("/api/revision")
@token_required
def get_revision():
    now = dt.datetime.now(dt.timezone.utc)
    rows = list(db.revision_items.find(owner_filter()).sort("due_at", 1))
    today, upcoming, overdue = [], [], []
    for r in rows:
        item = {"id": str(r["_id"]), "topic": r["topic"], "subject": r.get("subject", r["topic"]),
                 "strength": r.get("strength", 50)}
        due = aware(r.get("due_at", now))
        days = (due.date() - now.date()).days
        if days < 0:
            item["due"] = due.date().isoformat()
            overdue.append(item)
        elif days == 0:
            today.append(item)
        else:
            item["due"] = due.date().isoformat()
            upcoming.append(item)
    return jsonify(today=today, upcoming=upcoming, overdue=overdue)


# =========================================================== Analytics ====

@app.get("/api/analytics")
@token_required
def analytics():
    chats = db.chats.count_documents(owner_filter())
    documents_studied = db.documents.count_documents(owner_filter())
    attempts = list(db.quiz_attempts.find(owner_filter()))
    quiz_avg = round(sum(a["percentage"] for a in attempts) / len(attempts)) if attempts else 0

    topic_scores = {}
    for a in attempts:
        for topic, s in a.get("topics", {}).items():
            acc = topic_scores.setdefault(topic, {"correct": 0, "total": 0})
            acc["correct"] += s["correct"]
            acc["total"] += s["total"]
    topic_progress = [{"topic": t, "progress": round(s["correct"] / s["total"] * 100)}
                       for t, s in topic_scores.items()]
    strong = [t["topic"] for t in topic_progress if t["progress"] >= 70]
    weak = [t["topic"] for t in topic_progress if t["progress"] < 70]

    now = dt.datetime.now(dt.timezone.utc)
    weekly = []
    for i in range(6, -1, -1):
        day = (now - dt.timedelta(days=i)).date()
        day_chats = db.chats.count_documents({**owner_filter(), "role": "user",
                                               "created_at": {"$gte": dt.datetime.combine(day, dt.time.min, dt.timezone.utc),
                                                               "$lt": dt.datetime.combine(day, dt.time.max, dt.timezone.utc)}})
        day_attempts = [a for a in attempts if aware(a["created_at"]).date() == day]
        day_quiz = round(sum(a["percentage"] for a in day_attempts) / len(day_attempts)) if day_attempts else 0
        weekly.append({"day": day.strftime("%a"), "minutes": day_chats * 5, "quiz": day_quiz})

    trend = []
    for i in range(5, -1, -1):
        wk_attempts = [a for a in attempts if (now - aware(a["created_at"])).days // 7 == i]
        trend.append({"week": f"W{6-i}",
                       "score": round(sum(a["percentage"] for a in wk_attempts) / len(wk_attempts)) if wk_attempts else 0})

    return jsonify(
        stats={"learningTime": f"{chats * 3}m", "quizAverage": f"{quiz_avg}%",
               "topicsCompleted": len(strong), "studyStreak": g.user.get("streak", 0),
               "documentsStudied": documents_studied, "aiInteractions": chats},
        weeklyActivity=weekly, topicProgress=topic_progress, learningTrend=trend,
        strongTopics=strong, weakTopics=weak,
    )


@app.get("/api/recommendations")
@token_required
def recommendations():
    weak = list(db.revision_items.find(owner_filter()).sort("strength", 1).limit(4))
    if not weak:
        return jsonify(recommendations=[])
    out = []
    for w in weak:
        out.append({
            "id": str(w["_id"]), "title": f"Revise {w['topic']}",
            "reason": f"Your quiz strength on this topic is {w.get('strength', 50)}%.",
            "difficulty": "Hard" if w.get("strength", 50) < 40 else "Medium",
            "minutes": 30, "subject": w.get("subject", w["topic"]),
        })
    return jsonify(recommendations=out)


# ========================================================= Dashboard =====

@app.get("/api/learner/dashboard")
@token_required
def learner_dashboard():
    chats = db.chats.count_documents(owner_filter())
    attempts = list(db.quiz_attempts.find(owner_filter()))
    quiz_avg = round(sum(a["percentage"] for a in attempts) / len(attempts)) if attempts else 0
    documents_studied = db.documents.count_documents(owner_filter())
    revision_due = db.revision_items.count_documents(owner_filter())
    stats = [
        {"id": "time", "label": "Learning Time", "value": f"{chats*3}m", "delta": "", "trend": "up", "icon": "clock"},
        {"id": "chats", "label": "AI Chats", "value": str(chats), "delta": "", "trend": "up", "icon": "message"},
        {"id": "topics", "label": "Completed Topics", "value": str(documents_studied), "delta": "", "trend": "up", "icon": "check"},
        {"id": "quiz", "label": "Quiz Score", "value": f"{quiz_avg}%", "delta": "", "trend": "up", "icon": "target"},
        {"id": "progress", "label": "Study Progress",
         "value": f"{max(0, 100 - revision_due*10)}%", "delta": "", "trend": "up", "icon": "trending"},
    ]
    recent_docs = list(db.documents.find(owner_filter()).sort("uploaded_at", -1).limit(3))
    plan = list(db.study_plan.find(owner_filter()).sort("created_at", -1).limit(4))
    return jsonify(
        stats=stats,
        recentDocuments=[doc_id(r, {"uploadedAt": r["uploaded_at"].date().isoformat()}) for r in recent_docs],
        todayPlan=[doc_id(t) for t in plan],
        streak=g.user.get("streak", 0),
    )


# ============================================================= Profile ====

@app.get("/api/profile")
@token_required
def get_profile():
    attempts = list(db.quiz_attempts.find(owner_filter()))
    topic_scores = {}
    for a in attempts:
        for topic, s in a.get("topics", {}).items():
            acc = topic_scores.setdefault(topic, {"correct": 0, "total": 0})
            acc["correct"] += s["correct"]
            acc["total"] += s["total"]
    strong = [t for t, s in topic_scores.items() if s["correct"] / s["total"] >= 0.7]
    weak = [t for t, s in topic_scores.items() if s["correct"] / s["total"] < 0.7]
    return jsonify(profile={**public_user(g.user), "level": g.user.get("level", ""),
                             "goal": g.user.get("goal", ""), "streak": g.user.get("streak", 0)},
                   strongTopics=strong, weakTopics=weak)


@app.put("/api/profile")
@token_required
def update_profile():
    d = request.get_json(silent=True) or {}
    fields = {k: d[k] for k in ("name", "level", "goal") if k in d}
    if fields:
        db.users.update_one({"_id": g.user["_id"]}, {"$set": fields})
    row = db.users.find_one({"_id": g.user["_id"]})
    return jsonify(profile={**public_user(row), "level": row.get("level", ""), "goal": row.get("goal", "")})


# ========================================================== Institution ====

def inst_filter():
    return {"institution_id": g.user["_id"]}


@app.get("/api/institution/dashboard")
@institution_required
def institution_dashboard():
    total = db.students.count_documents(inst_filter())
    active = db.students.count_documents({**inst_filter(), "status": "active"})
    courses = list(db.courses.find(inst_filter()))
    avg_progress = round(sum(c.get("progress", 0) for c in courses) / len(courses)) if courses else 0
    students = list(db.students.find(inst_filter()))
    avg_quiz = round(sum(s.get("quizAverage", 0) for s in students) / len(students)) if students else 0
    stats = [
        {"id": "s1", "label": "Total Students", "value": str(total), "delta": "", "trend": "up"},
        {"id": "s2", "label": "Active Students", "value": str(active), "delta": "", "trend": "up"},
        {"id": "s3", "label": "Courses", "value": str(len(courses)), "delta": "", "trend": "up"},
        {"id": "s4", "label": "Average Progress", "value": f"{avg_progress}%", "delta": "", "trend": "up"},
        {"id": "s5", "label": "Average Quiz Score", "value": f"{avg_quiz}%", "delta": "", "trend": "up"},
    ]
    return jsonify(stats=stats, students=[doc_id(s) for s in students[:5]],
                    courses=[doc_id(c) for c in courses])


@app.get("/api/institution/students")
@institution_required
def list_students():
    q = request.args.get("q", "").strip().lower()
    rows = list(db.students.find(inst_filter()))
    if q:
        rows = [r for r in rows if q in r["name"].lower() or q in r["email"].lower()]
    return jsonify(students=[doc_id(r) for r in rows])


@app.post("/api/institution/students")
@institution_required
def add_student():
    d = request.get_json(silent=True) or {}
    s = {"institution_id": g.user["_id"], "name": d.get("name", ""), "email": d.get("email", ""),
         "course": d.get("course", ""), "progress": 0, "quizAverage": 0, "status": "active",
         "created_at": dt.datetime.now(dt.timezone.utc)}
    s["_id"] = db.students.insert_one(s).inserted_id
    return jsonify(student=doc_id(s)), 201


@app.patch("/api/institution/students/<id>")
@institution_required
def update_student(id):
    d = request.get_json(silent=True) or {}
    fields = {k: d[k] for k in ("status", "course", "name", "email") if k in d}
    db.students.update_one({"_id": oid_or_none(id), **inst_filter()}, {"$set": fields})
    row = db.students.find_one({"_id": oid_or_none(id), **inst_filter()})
    if not row:
        return jsonify(message="Student not found."), 404
    return jsonify(student=doc_id(row))


@app.delete("/api/institution/students/<id>")
@institution_required
def remove_student(id):
    db.students.delete_one({"_id": oid_or_none(id), **inst_filter()})
    return "", 204


@app.get("/api/institution/teachers")
@institution_required
def list_teachers():
    rows = db.teachers.find(inst_filter())
    return jsonify(teachers=[doc_id(r) for r in rows])


@app.post("/api/institution/teachers")
@institution_required
def add_teacher():
    d = request.get_json(silent=True) or {}
    t = {"institution_id": g.user["_id"], "name": d.get("name", ""), "email": d.get("email", ""),
         "department": d.get("department", ""), "courses": 0, "students": 0,
         "created_at": dt.datetime.now(dt.timezone.utc)}
    t["_id"] = db.teachers.insert_one(t).inserted_id
    return jsonify(teacher=doc_id(t)), 201


@app.delete("/api/institution/teachers/<id>")
@institution_required
def remove_teacher(id):
    db.teachers.delete_one({"_id": oid_or_none(id), **inst_filter()})
    return "", 204


@app.get("/api/institution/courses")
@institution_required
def list_courses():
    rows = db.courses.find(inst_filter())
    return jsonify(courses=[doc_id(r) for r in rows])


@app.post("/api/institution/courses")
@institution_required
def add_course():
    d = request.get_json(silent=True) or {}
    c = {"institution_id": g.user["_id"], "title": d.get("title", ""), "code": d.get("code", ""),
         "teacher": d.get("teacher", ""), "students": 0, "progress": 0,
         "created_at": dt.datetime.now(dt.timezone.utc)}
    c["_id"] = db.courses.insert_one(c).inserted_id
    return jsonify(course=doc_id(c)), 201


@app.delete("/api/institution/courses/<id>")
@institution_required
def remove_course(id):
    db.courses.delete_one({"_id": oid_or_none(id), **inst_filter()})
    return "", 204


@app.get("/api/institution/materials")
@institution_required
def list_materials():
    rows = db.materials.find(inst_filter())
    return jsonify(materials=[doc_id(r, {"uploadedAt": r["uploaded_at"].date().isoformat()}) for r in rows])


@app.post("/api/institution/materials")
@institution_required
def add_material():
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(message="No file received."), 400
    ext = (f.filename.rsplit(".", 1)[-1] if "." in f.filename else "").lower()
    ftype = "pdf" if ext == "pdf" else "docx" if ext in ("docx", "doc") else "pdf"
    raw = f.read()
    size_kb = len(raw) / 1024
    size = f"{size_kb/1024:.1f} MB" if size_kb > 1024 else f"{size_kb:.0f} KB"
    m = {"institution_id": g.user["_id"], "title": f.filename, "course": request.form.get("course", ""),
         "type": ftype, "size": size, "uploaded_at": dt.datetime.now(dt.timezone.utc)}
    m["_id"] = db.materials.insert_one(m).inserted_id
    return jsonify(material=doc_id(m, {"uploadedAt": m["uploaded_at"].date().isoformat()})), 201


@app.delete("/api/institution/materials/<id>")
@institution_required
def remove_material(id):
    db.materials.delete_one({"_id": oid_or_none(id), **inst_filter()})
    return "", 204


@app.get("/api/institution/progress")
@institution_required
def institution_progress():
    rows = db.students.find(inst_filter())
    return jsonify(students=[doc_id(r) for r in rows])


@app.get("/api/institution/analytics")
@institution_required
def institution_analytics():
    students = list(db.students.find(inst_filter()))
    courses = list(db.courses.find(inst_filter()))
    avg_progress = round(sum(s.get("progress", 0) for s in students) / len(students)) if students else 0
    avg_quiz = round(sum(s.get("quizAverage", 0) for s in students) / len(students)) if students else 0
    return jsonify(
        stats={"averageProgress": avg_progress, "averageQuiz": avg_quiz,
               "totalStudents": len(students), "totalCourses": len(courses)},
        courses=[doc_id(c) for c in courses],
    )


@app.get("/api/health")
def health():
    return jsonify(status="ok", model=os.getenv("LLM_PROVIDER", "ollama"))


if __name__ == "__main__":
    app.run(port=5000, debug=True)